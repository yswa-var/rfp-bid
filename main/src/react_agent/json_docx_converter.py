"""
JSON to DOCX Converter Library

Converts structured JSON data (configuration + content) into formatted .docx documents.
Supports headings, paragraphs, tables, images, headers, footers, and page management.

USAGE:
    Command Line:
        python json_docx_converter.py config.json content.json output.docx
    
    Programmatic:
        from json_docx_converter import convert_json_to_docx
        success, message = convert_json_to_docx('config.json', 'content.json', 'output.docx')
    
    Advanced (in-memory JSON):
        from json_docx_converter import ConfigParser, ContentParser, DocxGenerator
        config = ConfigParser(config_dict)
        content = ContentParser(content_dict)
        generator = DocxGenerator(config)
        generator.process_content(content)
        generator.save('output.docx')

FEATURES:
    - Headings (levels 1-6) with custom fonts, sizes, bold/italic
    - Paragraphs with custom fonts, sizes, alignment
    - Tables with header detection and styling
    - Images from file paths, URLs, or base64 encoding
    - Headers and footers with page numbers
    - Page settings (size, orientation, margins)
    - Automatic page breaks between sections (each section starts on a new page)
    - Manual page breaks with "new_page" directive
    - Style overrides per element

See example_config.json, example_content.json, and example_usage.py for detailed examples.
"""

import json
import base64
import io
import os
from typing import Dict, List, Any, Optional, Tuple
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import requests
from PIL import Image


class ConfigParser:
    """
    Parses and validates configuration JSON for document metadata and styling.
    Provides defaults for missing fields and validates required structure.
    """
    
    # Default configuration values
    DEFAULTS = {
        'metadata': {
            'title': 'Untitled Document',
            'author': 'Unknown',
            'page_size': 'A4',
            'orientation': 'portrait',
            'margins': {'top': 2.54, 'bottom': 2.54, 'left': 2.54, 'right': 2.54}
        },
        'styles': {
            'heading': [
                {'level': 1, 'font': 'Times New Roman', 'size': 16, 'bold': True},
                {'level': 2, 'font': 'Arial', 'size': 14, 'bold': True},
                {'level': 3, 'font': 'Arial', 'size': 12, 'italic': True}
            ],
            'paragraph': {'font': 'Courier New', 'size': 12},
            'table': {'border': True, 'cell_padding': 0.2, 'font': 'Arial', 'size': 10}
        }
    }
    
    def __init__(self, config_data: Dict[str, Any]):
        """
        Initialize ConfigParser with configuration data.
        
        Args:
            config_data: Dictionary containing configuration JSON
        """
        self.config_data = config_data
        self.metadata = self._parse_metadata()
        self.styles = self._parse_styles()
        self.header = self._parse_header_footer('header')
        self.footer = self._parse_header_footer('footer')
    
    def _parse_metadata(self) -> Dict[str, Any]:
        """Extract and validate metadata with defaults."""
        metadata = self.config_data.get('metadata', {})
        
        # Apply defaults for missing fields
        result = {
            'title': metadata.get('title', self.DEFAULTS['metadata']['title']),
            'author': metadata.get('author', self.DEFAULTS['metadata']['author']),
            'page_size': metadata.get('page_size', self.DEFAULTS['metadata']['page_size']),
            'orientation': metadata.get('orientation', self.DEFAULTS['metadata']['orientation']),
            'margins': metadata.get('margins', self.DEFAULTS['metadata']['margins'].copy())
        }
        
        # Validate orientation
        if result['orientation'] not in ['portrait', 'landscape']:
            raise ValueError(f"Invalid orientation: {result['orientation']}. Must be 'portrait' or 'landscape'.")
        
        return result
    
    def _parse_styles(self) -> Dict[str, Any]:
        """Extract and validate style definitions with defaults."""
        styles = self.config_data.get('styles', {})
        
        result = {
            'heading': [],
            'paragraph': styles.get('paragraph', self.DEFAULTS['styles']['paragraph'].copy()),
            'table': styles.get('table', self.DEFAULTS['styles']['table'].copy())
        }
        
        # Parse heading styles (support levels 1-6)
        heading_styles = styles.get('heading', self.DEFAULTS['styles']['heading'])
        for heading_def in heading_styles:
            level = heading_def.get('level', 1)
            if level < 1 or level > 6:
                raise ValueError(f"Invalid heading level: {level}. Must be between 1 and 6.")
            result['heading'].append({
                'level': level,
                'font': heading_def.get('font', 'Arial'),
                'size': heading_def.get('size', 12),
                'bold': heading_def.get('bold', False),
                'italic': heading_def.get('italic', False)
            })
        
        # Fill missing heading levels with defaults
        existing_levels = {h['level'] for h in result['heading']}
        for level in range(1, 7):
            if level not in existing_levels:
                result['heading'].append({
                    'level': level,
                    'font': 'Arial',
                    'size': max(12, 18 - level),
                    'bold': level <= 2,
                    'italic': False
                })
        
        # Sort heading styles by level
        result['heading'].sort(key=lambda x: x['level'])
        
        return result
    
    def _parse_header_footer(self, section_type: str) -> Optional[Dict[str, Any]]:
        """Parse header or footer configuration."""
        section = self.config_data.get(section_type)
        if not section:
            return None
        
        return {
            'text': section.get('text', ''),
            'align': section.get('align', 'left'),
            'font': section.get('font', 'Arial'),
            'size': section.get('size', 12),
            'image': section.get('image')
        }
    
    def get_heading_style(self, level: int) -> Dict[str, Any]:
        """Get style definition for a specific heading level."""
        for style in self.styles['heading']:
            if style['level'] == level:
                return style
        # Fallback if level not found
        return {'font': 'Arial', 'size': 12, 'bold': True, 'italic': False}


class ContentParser:
    """
    Parses and validates content JSON for document structure.
    Extracts elements by type and validates element-specific fields.
    """
    
    def __init__(self, content_data: Dict[str, List[Dict[str, Any]]]):
        """
        Initialize ContentParser with content data.
        
        Args:
            content_data: Dictionary with section names as keys, element lists as values
        """
        self.content_data = content_data
        self.sections = self._parse_sections()
    
    def _parse_sections(self) -> List[Tuple[str, List[Dict[str, Any]]]]:
        """
        Parse all sections and their elements.
        
        Returns:
            List of (section_name, elements) tuples
        """
        sections = []
        for section_name, elements in self.content_data.items():
            if not isinstance(elements, list):
                raise ValueError(f"Section '{section_name}' must contain a list of elements.")
            
            validated_elements = []
            for idx, element in enumerate(elements):
                validated = self._validate_element(element, section_name, idx)
                validated_elements.append(validated)
            
            sections.append((section_name, validated_elements))
        
        return sections
    
    def _validate_element(self, element: Dict[str, Any], section_name: str, idx: int) -> Dict[str, Any]:
        """
        Validate an individual element and ensure required fields exist.
        
        Args:
            element: Element dictionary
            section_name: Name of the section containing this element
            idx: Index of element in section (for error messages)
            
        Returns:
            Validated element dictionary
        """
        elem_type = element.get('type')
        if not elem_type:
            raise ValueError(f"Element {idx} in section '{section_name}' missing 'type' field.")
        
        # Validate based on element type
        if elem_type == 'heading':
            if 'text' not in element:
                raise ValueError(f"Heading element {idx} in '{section_name}' missing 'text' field.")
            if 'level' not in element:
                raise ValueError(f"Heading element {idx} in '{section_name}' missing 'level' field.")
            if element['level'] < 1 or element['level'] > 6:
                raise ValueError(f"Heading level must be 1-6, got {element['level']} in '{section_name}'.")
        
        elif elem_type == 'paragraph':
            if 'text' not in element:
                raise ValueError(f"Paragraph element {idx} in '{section_name}' missing 'text' field.")
        
        elif elem_type == 'table':
            if 'data' not in element:
                raise ValueError(f"Table element {idx} in '{section_name}' missing 'data' field.")
            if not isinstance(element['data'], list) or len(element['data']) == 0:
                raise ValueError(f"Table data must be a non-empty list in '{section_name}'.")
        
        elif elem_type == 'image':
            if 'path' not in element and 'url' not in element and 'base64' not in element:
                raise ValueError(f"Image element {idx} in '{section_name}' requires 'path', 'url', or 'base64'.")
        
        elif elem_type == 'new_page':
            # No validation needed for page breaks
            pass
        
        else:
            raise ValueError(f"Unknown element type '{elem_type}' in section '{section_name}'.")
        
        return element


class DocxGenerator:
    """
    Generates DOCX documents from parsed configuration and content.
    Handles all document elements: headings, paragraphs, tables, images, page breaks.
    """
    
    def __init__(self, config: ConfigParser):
        """
        Initialize DocxGenerator with configuration.
        
        Args:
            config: Parsed ConfigParser instance
        """
        self.config = config
        self.document = Document()
        self._apply_metadata()
        self._apply_page_settings()
        self._apply_headers_footers()
    
    def _apply_metadata(self):
        """Apply document metadata (title, author)."""
        core_props = self.document.core_properties
        core_props.title = self.config.metadata['title']
        core_props.author = self.config.metadata['author']
    
    def _apply_page_settings(self):
        """Apply page size, orientation, and margins to the document."""
        section = self.document.sections[0]
        
        # Set page size (A4 default: 21cm x 29.7cm, Letter: 8.5in x 11in)
        page_size = self.config.metadata['page_size'].upper()
        if page_size == 'A4':
            if self.config.metadata['orientation'] == 'portrait':
                section.page_width = Cm(21)
                section.page_height = Cm(29.7)
            else:
                section.page_width = Cm(29.7)
                section.page_height = Cm(21)
        elif page_size == 'LETTER':
            if self.config.metadata['orientation'] == 'portrait':
                section.page_width = Inches(8.5)
                section.page_height = Inches(11)
            else:
                section.page_width = Inches(11)
                section.page_height = Inches(8.5)
        
        # Set margins (provided in cm)
        margins = self.config.metadata['margins']
        section.top_margin = Cm(margins.get('top', 2.54))
        section.bottom_margin = Cm(margins.get('bottom', 2.54))
        section.left_margin = Cm(margins.get('left', 2.54))
        section.right_margin = Cm(margins.get('right', 2.54))
    
    def _apply_headers_footers(self):
        """Apply header and footer text with styling."""
        section = self.document.sections[0]
        
        # Apply header
        if self.config.header:
            header = section.header
            header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            header_para.text = self.config.header['text']
            header_para.alignment = self._get_alignment(self.config.header.get('align', 'left'))
            
            # Apply header font styling
            run = header_para.runs[0] if header_para.runs else header_para.add_run(self.config.header['text'])
            run.font.name = self.config.header.get('font', 'Arial')
            run.font.size = Pt(self.config.header.get('size', 12))
        
        # Apply footer (support page number placeholder)
        if self.config.footer:
            footer = section.footer
            footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            
            footer_text = self.config.footer['text']
            footer_para.alignment = self._get_alignment(self.config.footer.get('align', 'left'))
            
            # Handle page number placeholder
            if '{page_number}' in footer_text:
                parts = footer_text.split('{page_number}')
                if parts[0]:
                    run = footer_para.add_run(parts[0])
                    run.font.name = self.config.footer.get('font', 'Arial')
                    run.font.size = Pt(self.config.footer.get('size', 12))
                
                # Add page number field
                self._add_page_number(footer_para)
                
                if len(parts) > 1 and parts[1]:
                    run = footer_para.add_run(parts[1])
                    run.font.name = self.config.footer.get('font', 'Arial')
                    run.font.size = Pt(self.config.footer.get('size', 12))
            else:
                footer_para.text = footer_text
                run = footer_para.runs[0] if footer_para.runs else footer_para.add_run(footer_text)
                run.font.name = self.config.footer.get('font', 'Arial')
                run.font.size = Pt(self.config.footer.get('size', 12))
    
    def _add_page_number(self, paragraph):
        """Add a page number field to a paragraph."""
        run = paragraph.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = 'PAGE'
        
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)
    
    def _get_alignment(self, align: str) -> WD_ALIGN_PARAGRAPH:
        """Convert alignment string to WD_ALIGN_PARAGRAPH enum."""
        alignments = {
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
            'justify': WD_ALIGN_PARAGRAPH.JUSTIFY
        }
        return alignments.get(align.lower(), WD_ALIGN_PARAGRAPH.LEFT)
    
    def process_content(self, content: ContentParser):
        """
        Process all content sections and add elements to the document.
        Each section starts on a new page (except the first section).
        
        Args:
            content: Parsed ContentParser instance
        """
        for section_idx, (section_name, elements) in enumerate(content.sections):
            # Add page break before each section except the first
            if section_idx > 0:
                self._add_page_break()
            
            for element in elements:
                self._add_element(element)
    
    def _add_element(self, element: Dict[str, Any]):
        """
        Add a single element to the document based on its type.
        
        Args:
            element: Element dictionary with 'type' and type-specific fields
        """
        elem_type = element['type']
        
        if elem_type == 'heading':
            self._add_heading(element)
        elif elem_type == 'paragraph':
            self._add_paragraph(element)
        elif elem_type == 'table':
            self._add_table(element)
        elif elem_type == 'image':
            self._add_image(element)
        elif elem_type == 'new_page':
            self._add_page_break()
    
    def _add_heading(self, element: Dict[str, Any]):
        """Add a heading with configured styling."""
        level = element['level']
        text = element['text']
        
        # Get style from config or element override
        if 'font' in element or 'size' in element or 'bold' in element or 'italic' in element:
            # Element has style overrides
            style = {
                'font': element.get('font', self.config.get_heading_style(level)['font']),
                'size': element.get('size', self.config.get_heading_style(level)['size']),
                'bold': element.get('bold', self.config.get_heading_style(level)['bold']),
                'italic': element.get('italic', self.config.get_heading_style(level)['italic'])
            }
        else:
            # Use config style
            style = self.config.get_heading_style(level)
        
        # Add heading paragraph
        heading = self.document.add_heading(level=level)
        heading.text = text
        
        # Apply styling to the heading
        run = heading.runs[0] if heading.runs else heading.add_run(text)
        run.font.name = style['font']
        run.font.size = Pt(style['size'])
        run.font.bold = style['bold']
        run.font.italic = style['italic']
    
    def _add_paragraph(self, element: Dict[str, Any]):
        """Add a paragraph with configured styling."""
        text = element['text']
        style = self.config.styles['paragraph']
        
        # Element style overrides
        font = element.get('font', style['font'])
        size = element.get('size', style['size'])
        align = element.get('align', 'left')
        
        para = self.document.add_paragraph()
        para.alignment = self._get_alignment(align)
        
        run = para.add_run(text)
        run.font.name = font
        run.font.size = Pt(size)
    
    def _add_table(self, element: Dict[str, Any]):
        """
        Add a table from data array.
        Detects header rows (column_heading_* pattern) and applies styling.
        """
        data = element['data']
        if not data:
            return
        
        # Determine if first row is headers (contains column_heading_* keys)
        first_row = data[0]
        has_headers = any(key.startswith('column_heading_') for key in first_row.keys())
        
        if has_headers:
            # Extract header row
            headers = sorted([k for k in first_row.keys() if k.startswith('column_heading_')])
            header_values = [first_row[h] for h in headers]
            num_cols = len(headers)
            
            # Data rows start from index 1
            data_rows = data[1:]
        else:
            # All rows are data, infer columns from first row
            header_values = []
            data_rows = data
            num_cols = len(first_row.keys())
        
        # Create table
        num_rows = len(data_rows) + (1 if has_headers else 0)
        table = self.document.add_table(rows=num_rows, cols=num_cols)
        
        # Apply table style from config
        table_style = self.config.styles['table']
        if table_style.get('border', True):
            table.style = 'Table Grid'
        
        # Add header row if present
        if has_headers:
            header_row = table.rows[0]
            for col_idx, value in enumerate(header_values):
                cell = header_row.cells[col_idx]
                cell.text = str(value)
                # Apply header cell styling (bold)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.bold = True
                        run.font.name = table_style.get('font', 'Arial')
                        run.font.size = Pt(table_style.get('size', 10))
        
        # Add data rows
        row_offset = 1 if has_headers else 0
        for row_idx, row_data in enumerate(data_rows):
            table_row = table.rows[row_idx + row_offset]
            # Get values in order of keys
            values = [str(v) for v in row_data.values()]
            
            for col_idx, value in enumerate(values):
                if col_idx < num_cols:
                    cell = table_row.cells[col_idx]
                    cell.text = value
                    # Apply cell styling
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = table_style.get('font', 'Arial')
                            run.font.size = Pt(table_style.get('size', 10))
    
    def _add_image(self, element: Dict[str, Any]):
        """
        Add an image from file path, URL, or base64 data.
        Supports resizing and alignment.
        """
        try:
            image_stream = None
            
            # Load image from different sources
            if 'path' in element:
                # Load from file path
                image_path = element['path']
                if not os.path.exists(image_path):
                    print(f"Warning: Image file not found: {image_path}")
                    return
                image_stream = image_path
            
            elif 'url' in element:
                # Fetch from URL
                url = element['url']
                response = requests.get(url, timeout=10)
                response.raise_for_status()
                image_stream = io.BytesIO(response.content)
            
            elif 'base64' in element:
                # Decode base64
                base64_data = element['base64']
                image_data = base64.b64decode(base64_data)
                image_stream = io.BytesIO(image_data)
            
            if image_stream:
                # Add image with optional sizing
                width = element.get('width')
                height = element.get('height')
                
                # Create paragraph for alignment
                para = self.document.add_paragraph()
                align = element.get('align', 'left')
                para.alignment = self._get_alignment(align)
                
                # Add image to paragraph
                run = para.add_run()
                if width and height:
                    # Convert to inches (assuming pixels, approximate 96 DPI)
                    run.add_picture(image_stream, width=Inches(width/96), height=Inches(height/96))
                elif width:
                    run.add_picture(image_stream, width=Inches(width/96))
                elif height:
                    run.add_picture(image_stream, height=Inches(height/96))
                else:
                    run.add_picture(image_stream)
        
        except requests.RequestException as e:
            print(f"Warning: Failed to fetch image from URL: {e}")
        except Exception as e:
            print(f"Warning: Failed to add image: {e}")
    
    def _add_page_break(self):
        """Insert a page break."""
        self.document.add_page_break()
    
    def save(self, output_path: str):
        """
        Save the document to a file.
        
        Args:
            output_path: Path to save the .docx file
        """
        self.document.save(output_path)


def convert_json_to_docx(config_path: str, content_path: str, output_path: str) -> Tuple[bool, str]:
    """
    Main function to convert JSON files to a DOCX document.
    
    Args:
        config_path: Path to configuration JSON file
        content_path: Path to content JSON file
        output_path: Path to save the output .docx file
    
    Returns:
        Tuple of (success: bool, message: str)
    """
    try:
        # Load configuration JSON
        with open(config_path, 'r', encoding='utf-8') as f:
            config_data = json.load(f)
        
        # Load content JSON
        with open(content_path, 'r', encoding='utf-8') as f:
            content_data = json.load(f)
        
        # Parse configuration
        config = ConfigParser(config_data)
        
        # Parse content
        content = ContentParser(content_data)
        
        # Generate document
        generator = DocxGenerator(config)
        generator.process_content(content)
        
        # Save document
        generator.save(output_path)
        
        return True, f"Successfully generated document: {output_path}"
    
    except FileNotFoundError as e:
        return False, f"File not found: {e}"
    except json.JSONDecodeError as e:
        return False, f"Invalid JSON: {e}"
    except ValueError as e:
        return False, f"Validation error: {e}"
    except Exception as e:
        return False, f"Error generating document: {e}"


if __name__ == "__main__":
    # Example usage
    import sys
    
    if len(sys.argv) != 4:
        print("Usage: python json_docx_converter.py <config.json> <content.json> <output.docx>")
        sys.exit(1)
    
    config_file = sys.argv[1]
    content_file = sys.argv[2]
    output_file = sys.argv[3]
    
    success, message = convert_json_to_docx(config_file, content_file, output_file)
    print(message)
    sys.exit(0 if success else 1)

